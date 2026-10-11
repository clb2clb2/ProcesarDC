import re
import unicodedata
import threading
import shutil
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk

from openpyxl import load_workbook
from PyPDF2 import PdfReader, PdfWriter

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
except ImportError:
    DND_FILES = None
    TkinterDnD = None


@dataclass
class PdfPageData:
    source_name: str
    page_number: int
    fecha: str
    importe: Decimal | None
    factura: str
    beneficiario: str
    concepto: str


@dataclass
class ExcelRowData:
    row_number: int
    jg: str
    fecha: str
    importe_total: Decimal | None
    importe_retencion: Decimal | None
    factura: str
    beneficiario: str
    concepto: str

    @property
    def importe_neto(self) -> Decimal | None:
        if self.importe_total is None or self.importe_retencion is None:
            return None
        return (self.importe_total - self.importe_retencion).quantize(Decimal("0.01"))


@dataclass
class OriginSection:
    source_path: Path
    title: str
    start_page: int | None
    end_page: int | None
    is_numeric: bool
    numeric_id: str
    category: str
    from_marker: bool


BaseTk = TkinterDnD.Tk if TkinterDnD else tk.Tk


class App(BaseTk):
    def __init__(self) -> None:
        super().__init__()
        self.title("ProcesarDC")
        self.geometry("1500x860")
        self.minsize(1200, 740)

        self.base_dir = Path(__file__).resolve().parent
        self.ej_dir = self.base_dir / "ej"

        self.origin_pdf_paths: list[Path] = []
        self.jp_pdf_path: Path | None = None
        self.excel_path: Path | None = None

        self.jp_pdf_data: list[PdfPageData] = []
        self.excel_data: list[ExcelRowData] = []

        self.main_results: list[dict] = []
        self.excel_unassigned_results: list[dict] = []
        self.manual_jg_by_row_key: dict[str, str] = {}
        self.main_row_by_iid: dict[str, dict] = {}
        self.doc_type_var = tk.StringVar(value="Remesa INV")
        self.no_jp_mode_var = tk.BooleanVar(value=False)

        self.active_job: str | None = None

        self._build_ui()
        self._refresh_meta_fields()
        self._update_action_buttons()

    def _build_ui(self) -> None:
        root = ttk.Frame(self, padding=10)
        root.pack(fill=tk.BOTH, expand=True)

        top_buttons = ttk.Frame(root)
        top_buttons.pack(fill=tk.X)
        for col in range(3):
            top_buttons.columnconfigure(col, weight=1, uniform="top_thirds")

        self.btn_load_origin = ttk.Button(
            top_buttons,
            text="Cargar PDFs origen",
            command=self.load_origin_pdfs,
        )

        self.chk_no_jp_mode = ttk.Checkbutton(
            top_buttons,
            text="DC sin justificantes de pago (tarjetas o pago manual)",
            variable=self.no_jp_mode_var,
            command=self._on_toggle_no_jp_mode,
        )
        self.chk_no_jp_mode.grid(row=0, column=1, sticky="w", padx=3, pady=(0, 4))

        self.btn_load_origin.grid(row=1, column=0, sticky="ew", padx=(0, 6))

        self.btn_load_jp = ttk.Button(
            top_buttons,
            text="Cargar PDF Justif. Pago",
            command=self.load_jp_pdf,
        )
        self.btn_load_jp.grid(row=1, column=1, sticky="ew", padx=3)

        self.btn_load_excel = ttk.Button(
            top_buttons,
            text="Cargar Excel datos Justif. Gastos",
            command=self.load_excel,
        )
        self.btn_load_excel.grid(row=1, column=2, sticky="ew", padx=(6, 0))

        drop_row = ttk.Frame(root, padding=(0, 8, 0, 0))
        drop_row.pack(fill=tk.X)
        for col in range(3):
            drop_row.columnconfigure(col, weight=1, uniform="drop_thirds")

        self.lbl_drop_origin = self._build_drop_zone(
            drop_row,
            "Arrastra aqui los PDF de DC, CJ, JG, etc.",
            0,
        )
        self.lbl_drop_jp = self._build_drop_zone(
            drop_row,
            "Arrastra aqui el PDF de Justif. Pago",
            1,
        )
        self.lbl_drop_excel = self._build_drop_zone(
            drop_row,
            "Arrastra aqui el Excel de Justif. Gastos",
            2,
        )

        if DND_FILES:
            self._wire_drop(self.lbl_drop_origin, self._on_drop_origin)
            self._wire_drop(self.lbl_drop_jp, self._on_drop_jp)
            self._wire_drop(self.lbl_drop_excel, self._on_drop_excel)

        fields = ttk.LabelFrame(root, text="Datos detectados", padding=8)
        fields.pack(fill=tk.X, pady=(10, 8))

        self.field_vars = {
            "anio": tk.StringVar(value=""),
            "dc": tk.StringVar(value=""),
            "cj": tk.StringVar(value=""),
            "jg_origen": tk.StringVar(value="X"),
            "rt": tk.StringVar(value="—"),
            "otros": tk.StringVar(value="—"),
        }

        vcmd_numbers = (self.register(self._validate_numeric), "%P")

        self.entry_map: dict[str, tk.Entry] = {}

        self._add_field(fields, "AÑO", "anio", 0, editable=False)
        self._add_field(fields, "DC", "dc", 1, validate=vcmd_numbers)
        self._add_field(fields, "CJ", "cj", 2, validate=vcmd_numbers)
        self._add_field(fields, "JG", "jg_origen", 3, editable=False)
        self._add_field(fields, "RT", "rt", 4, validate=vcmd_numbers)
        self._add_field(fields, "OTROS", "otros", 5, editable=False)

        self.field_vars["dc"].trace_add("write", lambda *_: self._refresh_special_field_style("dc"))
        self.field_vars["cj"].trace_add("write", lambda *_: self._refresh_special_field_style("cj"))
        self.field_vars["jg_origen"].trace_add(
            "write", lambda *_: self._refresh_special_field_style("jg_origen")
        )

        actions = ttk.Frame(root)
        actions.pack(fill=tk.X, pady=(0, 8))
        actions.columnconfigure(0, weight=0)
        actions.columnconfigure(1, weight=1)
        actions.columnconfigure(2, weight=0)

        self.btn_analyze = ttk.Button(actions, text="Analizar", command=self.analyze)
        self.btn_analyze.grid(row=0, column=0, sticky="w")

        self.btn_reset = ttk.Button(actions, text="Reiniciar", command=self.confirm_reset_state)
        self.btn_reset.grid(row=0, column=2, sticky="e")

        self.tabs = ttk.Notebook(root)
        self.tabs.pack(fill=tk.BOTH, expand=True)

        tab_main = ttk.Frame(self.tabs)
        tab_unassigned_excel = ttk.Frame(self.tabs)
        self.tab_unassigned_excel = tab_unassigned_excel

        self.tabs.add(tab_main, text="Resultado del analisis")
        self.tabs.add(tab_unassigned_excel, text="Filas Excel no asignadas")

        columns = [
            ("page", "Pagina", 70),
            ("fecha", "Fecha operacion", 110),
            ("importe", "Importe ordenado", 130),
            ("factura", "Factura", 140),
            ("beneficiario", "Beneficiario", 280),
            ("concepto", "Concepto", 340),
            ("jg", "JG asignado", 150),
            ("jg_apply", "", 34),
            ("jg_candidates", "JG candidatos", 220),
        ]

        self.tree_main = self._build_table(tab_main, columns)
        self.tree_main.column("jg_apply", anchor="center", width=34, stretch=False)
        self.tree_main.tag_configure("incidencia", foreground="red")
        self.tree_main.tag_configure("probable_ok", foreground="green")
        self.tree_main.tag_configure("jg_duplicado", foreground="white", background="red")
        self.tree_main.bind("<Double-1>", self._on_main_double_click)
        self.tree_main.bind("<<TreeviewSelect>>", self._on_main_selection_changed)
        self.tree_main.bind("<ButtonRelease-1>", self._on_main_click_action)

        self.tree_unassigned_excel = self._build_table(
            tab_unassigned_excel,
            [
                ("jg", "JG", 120),
                ("fecha", "Fecha pago", 110),
                ("importe_neto", "Importe neto", 120),
                ("factura", "Factura", 140),
                ("beneficiario", "Beneficiario", 280),
                ("concepto", "Concepto", 560),
            ],
        )

        detail_frame = ttk.LabelFrame(tab_main, text="Detalle de discordancias", padding=(8, 6))
        detail_frame.pack(fill=tk.X, pady=(8, 0))

        self.detail_var = tk.StringVar(
            value="Seleccione una fila en rojo para ver el detalle de diferencias (JP/JG)."
        )
        self.lbl_mismatch_detail = ttk.Label(
            detail_frame,
            textvariable=self.detail_var,
            justify="left",
            anchor="w",
            wraplength=1420,
        )
        self.lbl_mismatch_detail.pack(fill=tk.X)

        process_frame = ttk.Frame(tab_main, padding=(0, 8, 0, 0))
        process_frame.pack(fill=tk.X)

        ttk.Label(process_frame, text="Tipo de Documento Contable:").pack(side=tk.LEFT)
        self.cmb_doc_type = ttk.Combobox(
            process_frame,
            textvariable=self.doc_type_var,
            values=["Remesa INV", "Tarjeta INV", "Remesa ORD", "Manual ORD", "Manual INV"],
            state="readonly",
            width=18,
        )
        self.cmb_doc_type.pack(side=tk.LEFT, padx=(8, 8))
        self.cmb_doc_type.current(0)

        self.btn_process = ttk.Button(process_frame, text="Procesar", command=self.process_documents)
        self.btn_process.pack(side=tk.LEFT)
        self._update_unassigned_tab_title()

        self.status_var = tk.StringVar(value="Cargue PDF origen, JP y Excel para comenzar.")
        ttk.Label(root, textvariable=self.status_var, padding=(0, 8, 0, 0)).pack(fill=tk.X)

        if not DND_FILES:
            self.status_var.set(
                "Drag&drop no disponible: instale tkinterdnd2. Puede usar los botones de carga."
            )

    def _build_drop_zone(self, parent: ttk.Frame, title: str, col: int) -> tk.Label:
        label = tk.Label(
            parent,
            text=title,
            relief="groove",
            bd=2,
            padx=10,
            pady=14,
            justify="center",
            anchor="center",
            bg="#f2f2f2",
        )
        label.grid(row=0, column=col, sticky="ew", padx=(0 if col == 0 else 3, 0 if col == 2 else 3))
        return label

    @staticmethod
    def _wire_drop(widget: tk.Label, handler) -> None:
        widget.drop_target_register(DND_FILES)
        widget.dnd_bind("<<Drop>>", handler)

    def _add_field(
        self,
        parent: ttk.LabelFrame,
        label: str,
        key: str,
        col: int,
        editable: bool = True,
        validate=None,
    ) -> None:
        parent.columnconfigure(col * 2 + 1, weight=1)

        ttk.Label(parent, text=label).grid(row=0, column=col * 2, sticky="w", padx=(0, 4))

        state = "normal" if editable else "readonly"
        entry = tk.Entry(
            parent,
            textvariable=self.field_vars[key],
            width=14,
            state=state,
            relief="solid",
            bd=1,
        )
        if editable and validate:
            entry.configure(validate="key", validatecommand=validate)

        entry.grid(row=0, column=col * 2 + 1, sticky="ew", padx=(0, 10))
        self.entry_map[key] = entry

    @staticmethod
    def _validate_numeric(proposed: str) -> bool:
        return proposed == "" or proposed.isdigit()

    @staticmethod
    def _validate_numeric_or_dash(proposed: str) -> bool:
        return proposed == "" or proposed.isdigit()

    def _refresh_special_field_style(self, field_key: str) -> None:
        entry = self.entry_map.get(field_key)
        if not entry:
            return

        value = self.field_vars[field_key].get().strip().upper()
        if value == "X":
            entry.configure(fg="red", font=("TkDefaultFont", 10, "bold"))
        else:
            entry.configure(fg="black", font=("TkDefaultFont", 10, "normal"))

    def _build_table(self, parent: ttk.Frame, columns: list[tuple[str, str, int]]) -> ttk.Treeview:
        container = ttk.Frame(parent)
        container.pack(fill=tk.BOTH, expand=True)

        col_ids = [c[0] for c in columns]
        tree = ttk.Treeview(container, columns=col_ids, show="headings")

        for col_id, title, width in columns:
            tree.heading(col_id, text=title)
            tree.column(col_id, width=width, anchor="w")

        y_scroll = ttk.Scrollbar(container, orient=tk.VERTICAL, command=tree.yview)
        x_scroll = ttk.Scrollbar(container, orient=tk.HORIZONTAL, command=tree.xview)
        tree.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)

        tree.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")

        container.rowconfigure(0, weight=1)
        container.columnconfigure(0, weight=1)
        return tree

    def _parse_drop_event_paths(self, raw_data: str) -> list[Path]:
        if not raw_data:
            return []

        paths: list[Path] = []
        for token in self.tk.splitlist(raw_data):
            clean = str(token).strip().strip("{}")
            if clean:
                paths.append(Path(clean))
        return paths

    def _on_drop_origin(self, event) -> None:
        files = [p for p in self._parse_drop_event_paths(event.data) if p.suffix.lower() == ".pdf"]
        if not files:
            messagebox.showwarning("Aviso", "Arrastre uno o mas archivos PDF validos.")
            return

        self._set_origin_pdfs(files)

    def _on_drop_jp(self, event) -> None:
        files = [p for p in self._parse_drop_event_paths(event.data) if p.suffix.lower() == ".pdf"]
        if len(files) != 1:
            messagebox.showwarning("Aviso", "Arrastre un unico PDF de justificantes de pago.")
            return

        self._set_jp_pdf(files[0])

    def _on_drop_excel(self, event) -> None:
        allowed = {".xlsx", ".xlsm", ".xltx", ".xltm"}
        files = [p for p in self._parse_drop_event_paths(event.data) if p.suffix.lower() in allowed]
        if len(files) != 1:
            messagebox.showwarning("Aviso", "Arrastre un unico archivo Excel valido.")
            return

        self._set_excel(files[0])

    def _update_action_buttons(self) -> None:
        if self.active_job is not None:
            self.btn_analyze.configure(state=tk.DISABLED)
            self.btn_process.configure(state=tk.DISABLED)
            self.btn_reset.configure(state=tk.DISABLED)
            return

        no_jp_mode = bool(self.no_jp_mode_var.get())
        can_analyze = not no_jp_mode and self.jp_pdf_path is not None and self.excel_path is not None
        can_process_without_analyze = no_jp_mode and len(self.origin_pdf_paths) > 0
        self.btn_analyze.configure(state=tk.NORMAL if can_analyze else tk.DISABLED)
        self.btn_process.configure(
            state=tk.NORMAL if (len(self.main_results) > 0 or can_process_without_analyze) else tk.DISABLED
        )
        self.btn_reset.configure(state=tk.NORMAL)

    def _on_toggle_no_jp_mode(self) -> None:
        no_jp_mode = bool(self.no_jp_mode_var.get())
        self.btn_load_jp.configure(state=tk.DISABLED if no_jp_mode else tk.NORMAL)
        self.btn_load_excel.configure(state=tk.DISABLED if no_jp_mode else tk.NORMAL)

        if no_jp_mode:
            self.status_var.set(
                "Modo sin justificantes de pago activo. No se requeriran JP ni Excel para procesar."
            )
        else:
            self.status_var.set("Modo normal activo. Cargue PDF JP y Excel para analizar.")

        self._update_action_buttons()

    def load_origin_pdfs(self) -> None:
        paths = filedialog.askopenfilenames(
            title="Seleccionar PDFs origen",
            initialdir=str(self.ej_dir if self.ej_dir.exists() else self.base_dir),
            filetypes=[("PDF", "*.pdf")],
        )
        if not paths:
            return

        files = [Path(p) for p in paths]
        self._set_origin_pdfs(files)

    def _set_origin_pdfs(self, files: list[Path]) -> None:
        unique_sorted = sorted({f.resolve() for f in files if f.exists()}, key=lambda p: p.name.lower())
        if not unique_sorted:
            messagebox.showwarning("Aviso", "No se han seleccionado PDFs origen validos.")
            return

        error_message = self._validate_origin_jg_files(unique_sorted)
        if error_message:
            messagebox.showerror("Error", error_message)
            return

        self.origin_pdf_paths = list(unique_sorted)
        names = [f.stem for f in self.origin_pdf_paths]
        self.lbl_drop_origin.configure(text=", ".join(names) if names else "")
        self._refresh_meta_fields()

        self.status_var.set(f"PDFs origen cargados: {len(self.origin_pdf_paths)}")
        self._update_action_buttons()

    def load_jp_pdf(self) -> None:
        path = filedialog.askopenfilename(
            title="Seleccionar PDF Justif. Pago",
            initialdir=str(self.ej_dir if self.ej_dir.exists() else self.base_dir),
            filetypes=[("PDF", "*.pdf")],
        )
        if not path:
            return

        self._set_jp_pdf(Path(path))

    def _set_jp_pdf(self, path: Path) -> None:
        try:
            reader = PdfReader(str(path))
            page_count = len(reader.pages)
        except Exception as exc:
            messagebox.showerror("Error", f"No se pudo leer el PDF JP: {exc}")
            return

        self.jp_pdf_path = path
        self.lbl_drop_jp.configure(text=f"JP encontrados: {page_count}")
        self.status_var.set("PDF de justificantes de pago cargado correctamente.")
        self._update_action_buttons()

    def load_excel(self) -> None:
        path = filedialog.askopenfilename(
            title="Seleccionar Excel datos Justif. Gastos",
            initialdir=str(self.ej_dir if self.ej_dir.exists() else self.base_dir),
            filetypes=[("Excel", "*.xlsx *.xlsm *.xltx *.xltm")],
        )
        if not path:
            return

        self._set_excel(Path(path))

    def _set_excel(self, path: Path) -> None:
        try:
            excel_rows = self._read_excel(path)
        except Exception as exc:
            messagebox.showerror("Error", f"No se pudo leer el Excel: {exc}")
            return

        self.excel_path = path
        self.excel_data = excel_rows
        self.lbl_drop_excel.configure(text=f"JG en la tabla Excel: {len(self.excel_data)}")
        self._refresh_meta_fields()
        self.status_var.set("Excel de justificantes de gastos cargado correctamente.")
        self._update_action_buttons()

    def _refresh_meta_fields(self) -> None:
        sections, _ = self._scan_origin_sections(self.origin_pdf_paths)
        self._refresh_anio_field()
        self._refresh_prefixed_number_field("dc", "DC", missing_value="X", sections=sections)
        self._refresh_prefixed_number_field("cj", "CJ", missing_value="X", sections=sections)
        self._refresh_jg_markers_field(sections)
        self._refresh_prefixed_number_field("rt", "RT", missing_value="—", sections=sections)
        self._refresh_otros_field(sections)
        self._refresh_special_field_style("dc")
        self._refresh_special_field_style("cj")
        self._refresh_special_field_style("jg_origen")

    def _refresh_anio_field(self) -> None:
        years: set[str] = set()
        for row in self.excel_data:
            jg = str(row.jg or "").strip()
            match = re.match(r"^(20\d{2})/\d+$", jg)
            if match:
                years.add(match.group(1))

        if not years:
            self.field_vars["anio"].set("")
            return

        if len(years) == 1:
            self.field_vars["anio"].set(next(iter(years)))
            return

        self.field_vars["anio"].set("/".join(sorted(years)))

    def _refresh_prefixed_number_field(
        self,
        field_key: str,
        prefix: str,
        missing_value: str,
        sections: list[OriginSection],
    ) -> None:
        prefix_upper = prefix.upper()
        matched_numbers: list[str] = []
        for section in sections:
            if section.is_numeric:
                continue
            number = self._extract_prefixed_number(section.title, prefix_upper)
            if number:
                matched_numbers.append(number)

        if not matched_numbers:
            self.field_vars[field_key].set(missing_value)
            self.entry_map[field_key].configure(state="readonly")
            return

        self.entry_map[field_key].configure(state="normal")
        self.field_vars[field_key].set(matched_numbers[0])

    def _refresh_jg_markers_field(self, sections: list[OriginSection]) -> None:
        count = len({s.numeric_id for s in sections if s.is_numeric and s.numeric_id})
        if count <= 0:
            self.field_vars["jg_origen"].set("X")
            return

        self.field_vars["jg_origen"].set(f"JG encontrados: {count}")

    def _validate_origin_jg_files(self, files: list[Path]) -> str:
        sections, errors = self._scan_origin_sections(files)
        if errors:
            return errors[0]

        dc_sections = [s for s in sections if not s.is_numeric and s.category == "DC"]
        if len(dc_sections) > 1:
            return "¡Atención! Debes cargar un solo archivo de Documento Contable (DC)."

        cj_sections = [s for s in sections if not s.is_numeric and s.category == "CJ"]
        if len(cj_sections) > 1:
            return "¡Atención! Debes cargar un solo archivo de Cuenta Justificativa (CJ)."

        key_to_origin: dict[str, str] = {}
        for section in sections:
            if section.is_numeric:
                key = f"JG::{section.numeric_id}"
            else:
                key = f"TXT::{self._normalize_text(section.title)}"

            if key in key_to_origin:
                return (
                    "¡Atención! Hay marcadores o archivos en conflicto "
                    "(duplicados entre secciones extraídas y/o archivos directos)."
                )
            key_to_origin[key] = section.source_path.name

        return ""

    def _scan_origin_sections(self, files: list[Path]) -> tuple[list[OriginSection], list[str]]:
        sections: list[OriginSection] = []
        errors: list[str] = []

        for pdf_path in files:
            stem_upper = pdf_path.stem.strip().upper()
            if self._is_strict_prefixed_file_name(stem_upper):
                sections.append(
                    OriginSection(
                        source_path=pdf_path,
                        title=pdf_path.stem.strip(),
                        start_page=None,
                        end_page=None,
                        is_numeric=False,
                        numeric_id="",
                        category=self._classify_non_numeric_title(pdf_path.stem),
                        from_marker=False,
                    )
                )
                continue

            try:
                reader = PdfReader(str(pdf_path))
                outline = getattr(reader, "outline", None)
                if outline is None:
                    outline = getattr(reader, "outlines", None)
                entries = self._extract_outline_entries(reader, outline)
            except Exception as exc:
                errors.append(f"No se pudo leer el PDF de origen '{pdf_path.name}': {exc}")
                continue

            if not entries:
                title = pdf_path.stem.strip()
                if title:
                    sections.append(
                        OriginSection(
                            source_path=pdf_path,
                            title=title,
                            start_page=None,
                            end_page=None,
                            is_numeric=False,
                            numeric_id="",
                            category=self._classify_non_numeric_title(title),
                            from_marker=False,
                        )
                    )
                continue

            total_pages = len(reader.pages)
            entries_sorted = sorted(entries, key=lambda it: (it[1], it[2]))
            for idx, (title_raw, start_page, _order) in enumerate(entries_sorted):
                end_page = entries_sorted[idx + 1][1] if idx + 1 < len(entries_sorted) else total_pages
                if start_page < 0 or end_page > total_pages or start_page >= end_page:
                    continue

                title = self._normalize_section_title(title_raw)
                if not title:
                    continue

                is_numeric = title.isdigit()
                numeric_id = (title.lstrip("0") or "0") if is_numeric else ""
                category = "JG" if is_numeric else self._classify_non_numeric_title(title)
                sections.append(
                    OriginSection(
                        source_path=pdf_path,
                        title=title,
                        start_page=start_page,
                        end_page=end_page,
                        is_numeric=is_numeric,
                        numeric_id=numeric_id,
                        category=category,
                        from_marker=True,
                    )
                )

        return sections, errors

    def _extract_outline_entries(self, reader: PdfReader, outline) -> list[tuple[str, int, int]]:
        entries: list[tuple[str, int, int]] = []
        order_counter = 0

        def _walk(node) -> None:
            nonlocal order_counter
            if node is None:
                return
            if isinstance(node, list):
                for item in node:
                    _walk(item)
                return

            title = self._extract_outline_title(node)
            if title is not None:
                try:
                    page_number = int(reader.get_destination_page_number(node))
                    entries.append((str(title), page_number, order_counter))
                    order_counter += 1
                except Exception:
                    pass

        _walk(outline)
        return entries

    @staticmethod
    def _normalize_section_title(value: str) -> str:
        return re.sub(r"\s+", " ", str(value or "").strip())

    @staticmethod
    def _extract_prefixed_number(title: str, prefix: str) -> str:
        text = str(title or "").strip().upper()
        match = re.match(rf"^{prefix}(\d+)", text)
        if not match:
            return ""
        return match.group(1).lstrip("0") or "0"

    @staticmethod
    def _classify_non_numeric_title(title: str) -> str:
        text = str(title or "").strip().upper()
        if re.match(r"^DC\d+", text):
            return "DC"
        if re.match(r"^CJ\d+", text):
            return "CJ"
        if re.match(r"^RT\d+", text):
            return "RT"
        return "OTHER"

    @staticmethod
    def _is_strict_prefixed_file_name(stem: str) -> bool:
        text = str(stem or "").strip().upper()
        return bool(re.match(r"^(DC|CJ|RT)\d+", text))

    def _extract_numeric_bookmark_titles(self, outline) -> list[str]:
        titles = self._extract_outline_titles(outline)
        normalized: list[str] = []
        for title in titles:
            text = title.strip()
            if text.isdigit():
                normalized.append(text.lstrip("0") or "0")
        return normalized

    def _extract_numeric_bookmark_titles(self, outline) -> list[str]:
        titles = self._extract_outline_titles(outline)
        normalized: list[str] = []
        for title in titles:
            text = title.strip()
            if text.isdigit():
                normalized.append(text.lstrip("0") or "0")
        return normalized

    def _extract_outline_titles(self, outline) -> list[str]:
        collected: list[str] = []

        def _walk(node) -> None:
            if node is None:
                return
            if isinstance(node, list):
                for item in node:
                    _walk(item)
                return

            title_text = self._extract_outline_title(node)
            if title_text is not None:
                collected.append(title_text)

        _walk(outline)
        return collected

    @staticmethod
    def _extract_outline_title(outline) -> str | None:
        if hasattr(outline, "title"):
            value = getattr(outline, "title", None)
            return str(value) if value is not None else None

        if isinstance(outline, dict) and "/Title" in outline:
            value = outline.get("/Title")
            return str(value) if value is not None else None

        return None

    def _refresh_otros_field(self, sections: list[OriginSection]) -> None:
        others = [s.title for s in sections if not s.is_numeric and s.category == "OTHER"]
        self.field_vars["otros"].set(", ".join(others) if others else "—")

    def analyze(self) -> None:
        if self.no_jp_mode_var.get():
            messagebox.showinfo(
                "Analizar",
                "En modo sin justificantes de pago no es necesario analizar JP ni Excel.",
            )
            return

        if self.jp_pdf_path is None or self.excel_path is None:
            messagebox.showwarning("Aviso", "Debe cargar el PDF JP y el Excel antes de analizar.")
            return

        if self.active_job is not None:
            return

        self.active_job = "analyze"
        self._update_action_buttons()
        self.status_var.set("Analizando JP vs JG...")

        worker = threading.Thread(target=self._run_analyze_worker, daemon=True)
        worker.start()

    def process_documents(self) -> None:
        validation_error = self._validate_before_processing()
        if validation_error:
            messagebox.showerror("Error", validation_error)
            return

        parent_dir = filedialog.askdirectory(
            title="Seleccionar carpeta destino",
            initialdir=str(self.base_dir),
        )
        if not parent_dir:
            return

        try:
            year = self._resolve_output_year()
            if not year:
                return

            dc_number = int(self.field_vars["dc"].get().strip())
            doc_type_upper = self.doc_type_var.get().strip().upper()
            folder_name = f"{year}{dc_number:09d} ({doc_type_upper})"

            output_dir = Path(parent_dir) / folder_name
            output_dir.mkdir(parents=True, exist_ok=True)

            sections, errors = self._scan_origin_sections(self.origin_pdf_paths)
            if errors:
                raise ValueError(errors[0])

            copy_plan, extract_plan = self._build_output_plans_from_sections(output_dir, sections)

            all_targets = [dst for _, dst in copy_plan] + [dst for _, dst in extract_plan]
            duplicates = self._find_duplicate_paths(all_targets)
            if duplicates:
                messagebox.showerror(
                    "Error",
                    "Existen conflictos internos de nombres en los archivos a generar. Revise los datos.",
                )
                return

            existing_targets = [p for p in all_targets if p.exists()]
            if existing_targets:
                proceed = messagebox.askyesno(
                    "Conflicto de nombres",
                    (
                        "Ya existen archivos con el mismo nombre en la carpeta destino.\n"
                        "Si continua, se sobrescribiran.\n\n"
                        "Desea continuar?"
                    ),
                )
                if not proceed:
                    return

            for src, dst in copy_plan:
                shutil.copy2(src, dst)

            generated_jg_files = self._extract_sections_to_pdfs(extract_plan)
            if not self.no_jp_mode_var.get():
                self._append_jp_pages_to_generated_jg(output_dir, generated_jg_files)

            self.status_var.set(f"Proceso completado. Carpeta generada: {output_dir}")
            messagebox.showinfo("Procesar", "Proceso completado correctamente.")
        except Exception as exc:
            messagebox.showerror("Error", f"No se pudo completar el procesamiento: {exc}")

    def _resolve_output_year(self) -> str:
        current_year = str(datetime.now().year)
        existing_year = self.field_vars["anio"].get().strip()

        if re.match(r"^\d{4}$", existing_year):
            return existing_year

        if not self.no_jp_mode_var.get():
            return ""

        suggested = current_year
        chosen = simpledialog.askstring(
            "Confirmar año",
            "No se detectó año en los datos. Indique el año para la carpeta de salida:",
            initialvalue=suggested,
            parent=self,
        )

        if chosen is None:
            return ""

        chosen = str(chosen).strip()
        if not re.match(r"^\d{4}$", chosen):
            messagebox.showerror("Error", "El año indicado no es válido. Debe tener 4 dígitos.")
            return ""

        self.field_vars["anio"].set(chosen)
        return chosen

    def _validate_before_processing(self) -> str:
        no_jp_mode = bool(self.no_jp_mode_var.get())

        sections, errors = self._scan_origin_sections(self.origin_pdf_paths)
        if errors:
            return errors[0]

        dc_sections = [s for s in sections if not s.is_numeric and s.category == "DC"]
        if len(dc_sections) != 1:
            return "Debe haber exactamente un archivo o sección de Documento Contable (DC)"

        cj_sections = [s for s in sections if not s.is_numeric and s.category == "CJ"]
        if len(cj_sections) != 1:
            return "Debe haber exactamente un archivo o sección de Cuenta Justificativa (CJ)"

        if str(self.field_vars["dc"].get()).strip().upper() == "X":
            return "No se ha cargado un Documento Contable válido"

        if str(self.field_vars["cj"].get()).strip().upper() == "X":
            return "No se ha cargado un Cuenta Justificativa válida"

        if str(self.field_vars["jg_origen"].get()).strip().upper() == "X":
            return "No se han cargado Justificantes de Gasto válidos, o el archivo no tiene marcadores"

        if not self._is_positive_numeric_field("dc") or not self._is_positive_numeric_field("cj"):
            return "Debe escribir el número de Documento Contable/Cuenta Justificativa"

        if no_jp_mode:
            year = self.field_vars["anio"].get().strip()
            if year and not re.match(r"^\d{4}$", year):
                return "El año detectado no es válido"
            return ""

        if not self.main_results:
            return "Debe ejecutar el análisis antes de procesar"

        if any(not self._get_effective_jg(row) for row in self.main_results):
            return "Hay Justificantes de Pago a los que no se ha asignado un número de JG"

        numeric_assignments = [self._extract_numeric_jg_id(self._get_effective_jg(row)) for row in self.main_results]
        if any(not x for x in numeric_assignments):
            return "Hay Justificantes de Pago a los que no se ha asignado un número de JG"

        if len(set(numeric_assignments)) != len(numeric_assignments):
            return "Hay varias páginas de JP asignadas al mismo JG"

        marker_set = {s.numeric_id for s in sections if s.is_numeric and s.numeric_id}
        if not marker_set:
            return "No se han cargado Justificantes de Gasto válidos, o el archivo no tiene marcadores"

        if any(num not in marker_set for num in numeric_assignments):
            return "Hay Justificantes de Pago con JG asignado que no existe en el archivo JG"

        year = self.field_vars["anio"].get().strip()
        if not re.match(r"^\d{4}$", year):
            return "El año detectado no es válido"

        return ""

    def _is_positive_numeric_field(self, key: str) -> bool:
        value = str(self.field_vars[key].get()).strip()
        if not value.isdigit():
            return False
        return int(value) > 0

    @staticmethod
    def _find_duplicate_paths(paths: list[Path]) -> list[Path]:
        seen: dict[str, int] = {}
        duplicates: list[Path] = []
        for path in paths:
            key = str(path.resolve()).lower()
            seen[key] = seen.get(key, 0) + 1
            if seen[key] == 2:
                duplicates.append(path)
        return duplicates

    def _build_output_plans_from_sections(
        self,
        output_dir: Path,
        sections: list[OriginSection],
    ) -> tuple[list[tuple[Path, Path]], list[tuple[OriginSection, Path]]]:
        copy_plan: list[tuple[Path, Path]] = []
        extract_plan: list[tuple[OriginSection, Path]] = []
        dc_num = str(int(self.field_vars["dc"].get().strip()))
        cj_num = str(int(self.field_vars["cj"].get().strip()))

        for section in sections:
            if section.is_numeric:
                target_name = f"JG{section.numeric_id}.pdf"
            elif section.category == "DC":
                target_name = f"DC{dc_num}.pdf"
            elif section.category == "CJ":
                target_name = f"CJ{cj_num}.pdf"
            elif section.from_marker:
                target_name = f"{self._safe_output_stem(section.title)}.pdf"
            else:
                target_name = section.source_path.name

            target_path = output_dir / target_name
            if section.from_marker:
                extract_plan.append((section, target_path))
            else:
                copy_plan.append((section.source_path, target_path))

        return copy_plan, extract_plan

    def _extract_sections_to_pdfs(self, extract_plan: list[tuple[OriginSection, Path]]) -> dict[str, Path]:
        generated_jg_files: dict[str, Path] = {}
        readers_by_path: dict[str, PdfReader] = {}

        for section, out_file in extract_plan:
            if section.start_page is None or section.end_page is None:
                continue

            source_key = str(section.source_path.resolve())
            if source_key not in readers_by_path:
                readers_by_path[source_key] = PdfReader(str(section.source_path))

            reader = readers_by_path[source_key]
            writer = PdfWriter()
            for page_idx in range(section.start_page, section.end_page):
                writer.add_page(reader.pages[page_idx])

            with open(out_file, "wb") as fh:
                writer.write(fh)

            if section.is_numeric and section.numeric_id:
                generated_jg_files[section.numeric_id] = out_file

        return generated_jg_files

    @staticmethod
    def _safe_output_stem(value: str) -> str:
        stem = re.sub(r"[\\/:*?\"<>|]", "_", str(value or "").strip())
        stem = stem.rstrip(". ")
        return stem or "SECCION"

    def _append_jp_pages_to_generated_jg(self, output_dir: Path, generated_jg_files: dict[str, Path]) -> None:
        assert self.jp_pdf_path is not None
        jp_reader = PdfReader(str(self.jp_pdf_path))

        assignments: dict[str, list[int]] = defaultdict(list)
        ordered_rows = sorted(self.main_results, key=lambda r: int(r.get("page", 0)))
        for row in ordered_rows:
            page_number = int(row.get("page", 0))
            effective_jg = self._get_effective_jg(row)
            numeric_id = self._extract_numeric_jg_id(effective_jg)
            if page_number > 0 and numeric_id:
                assignments[numeric_id].append(page_number)

        for numeric_id, page_numbers in assignments.items():
            target_pdf = generated_jg_files.get(numeric_id)
            if target_pdf is None:
                raise ValueError(f"No existe archivo JG para el identificador {numeric_id}")

            existing_reader = PdfReader(str(target_pdf))
            writer = PdfWriter()

            for page in existing_reader.pages:
                writer.add_page(page)

            for page_number in page_numbers:
                jp_idx = page_number - 1
                if jp_idx < 0 or jp_idx >= len(jp_reader.pages):
                    raise ValueError(f"Pagina JP fuera de rango: {page_number}")
                writer.add_page(jp_reader.pages[jp_idx])

            with open(target_pdf, "wb") as fh:
                writer.write(fh)

    def _extract_numeric_jg_id(self, value: str) -> str:
        text = str(value or "").strip().upper()
        if not text:
            return ""

        match = re.match(r"^(\d+)", text)
        if not match:
            return ""

        return match.group(1).lstrip("0") or "0"

    def _run_analyze_worker(self) -> None:
        try:
            assert self.jp_pdf_path is not None
            assert self.excel_path is not None

            pdf_data = self._read_pdf_data(self.jp_pdf_path)
            excel_data = self._read_excel(self.excel_path)

            main_results: list[dict] = []
            for page in pdf_data:
                matched_rows = self._find_excel_matches(page, excel_data)
                probable_jg_raw, probable_jg_source = self._extract_probable_jg(page)
                probable_jg = self._jg_number_only(probable_jg_raw)
                row_key = self._build_row_key(page.source_name, page.page_number)

                row_base = {
                    "source": page.source_name,
                    "page": page.page_number,
                    "row_key": row_key,
                    "fecha": page.fecha,
                    "importe": self._format_decimal(page.importe),
                    "factura": page.factura,
                    "beneficiario": page.beneficiario,
                    "concepto": page.concepto,
                    "mismatch_details": [],
                }

                if len(matched_rows) == 1:
                    matched_jg = self._jg_number_only(self._format_jg(matched_rows[0].jg))
                    row = dict(row_base)
                    row["probable_jg"] = probable_jg
                    row["probable_jg_source"] = probable_jg_source
                    row["jg_candidates"] = matched_jg
                    row["match_status"] = "found"

                    if probable_jg and probable_jg == matched_jg:
                        row["jg_base"] = matched_jg
                        row["jg_display_base"] = matched_jg
                        row["jg_candidates"] = ""
                        row["row_color"] = "green"
                    elif probable_jg and probable_jg != matched_jg and probable_jg_source == "concepto":
                        row["jg_base"] = matched_jg
                        row["jg_display_base"] = matched_jg
                        row["jg_candidates"] = probable_jg
                        row["row_color"] = "red"
                    else:
                        row["jg_base"] = matched_jg
                        row["jg_display_base"] = matched_jg
                        row["row_color"] = ""

                    row["jg"] = row["jg_display_base"]
                    main_results.append(row)
                elif len(matched_rows) == 0:
                    near_candidates = self._find_near_excel_candidates(page, excel_data, max_mismatches=2)

                    row = dict(row_base)
                    row["probable_jg"] = probable_jg
                    row["probable_jg_source"] = probable_jg_source
                    row["jg_base"] = ""
                    row["jg_display_base"] = ""
                    row["jg"] = ""

                    if near_candidates:
                        best = near_candidates[0]
                        best_count = best["mismatches_count"]
                        best_group = [c for c in near_candidates if c["mismatches_count"] == best_count]

                        row["jg_candidates"] = " | ".join(
                            self._jg_number_only(self._format_jg(c["row"].jg)) for c in best_group
                        )
                        row["mismatch_details"] = best["mismatches"][:2]
                    else:
                        row["jg_candidates"] = ""

                    row["match_status"] = "not_found"
                    row["row_color"] = "red"
                    main_results.append(row)
                else:
                    candidates_list = [self._jg_number_only(self._format_jg(r.jg)) for r in matched_rows]
                    probable_in_matches = probable_jg and probable_jg in candidates_list

                    if probable_jg and not probable_in_matches:
                        candidates_list.append(probable_jg)

                    candidates = " | ".join(dict.fromkeys(candidates_list))

                    row = dict(row_base)
                    row["probable_jg"] = probable_jg
                    row["probable_jg_source"] = probable_jg_source
                    row["jg_base"] = probable_jg if probable_in_matches else ""
                    row["jg_display_base"] = probable_jg if probable_in_matches else ""
                    row["jg"] = row["jg_display_base"]
                    row["jg_candidates"] = candidates
                    row["match_status"] = "dubious"
                    row["row_color"] = "red"
                    main_results.append(row)

            self.after(0, lambda: self._on_analyze_done(pdf_data, excel_data, main_results))
        except Exception as exc:
            self.after(0, lambda: self._on_analyze_failed(str(exc)))

    def _on_analyze_done(
        self,
        pdf_data: list[PdfPageData],
        excel_data: list[ExcelRowData],
        main_results: list[dict],
    ) -> None:
        self.jp_pdf_data = list(pdf_data)
        self.excel_data = list(excel_data)
        self.main_results = list(main_results)
        self.manual_jg_by_row_key = {}
        self._rebuild_excel_unassigned_results()
        self._refresh_tables()

        self.active_job = None
        self._update_action_buttons()
        self._update_status_summary(prefix="Analisis completado. ")

    def _on_analyze_failed(self, message: str) -> None:
        self.active_job = None
        self._update_action_buttons()
        messagebox.showerror("Error", f"Fallo durante el analisis: {message}")

    @staticmethod
    def _build_row_key(source_name: str, page_number: int) -> str:
        return f"{source_name}::{page_number}"

    def _refresh_tables(self) -> None:
        for item in self.tree_main.get_children():
            self.tree_main.delete(item)
        self.main_row_by_iid = {}

        duplicated_jg = self._get_duplicated_assigned_jg_set()

        for row in self.main_results:
            values = [
                row.get("page", ""),
                row.get("fecha", ""),
                row.get("importe", ""),
                row.get("factura", ""),
                row.get("beneficiario", ""),
                row.get("concepto", ""),
                row.get("jg", ""),
                "<-" if str(row.get("jg_candidates", "")).strip() else "",
                row.get("jg_candidates", ""),
            ]
            row_jg = str(row.get("jg", "")).strip().upper()
            if row_jg and row_jg in duplicated_jg:
                tags = ("jg_duplicado",)
            elif row.get("row_color") == "green":
                tags = ("probable_ok",)
            elif row.get("row_color") == "red" or row.get("match_status") in {"not_found", "dubious"}:
                tags = ("incidencia",)
            else:
                tags = ()

            iid = self.tree_main.insert("", tk.END, values=values, tags=tags)
            self.main_row_by_iid[iid] = row

        self._fill_tree(
            self.tree_unassigned_excel,
            self.excel_unassigned_results,
            ["jg", "fecha", "importe_neto", "factura", "beneficiario", "concepto"],
        )
        self._update_unassigned_tab_title()

    @staticmethod
    def _fill_tree(tree: ttk.Treeview, rows: list[dict], keys: list[str]) -> None:
        for item in tree.get_children():
            tree.delete(item)

        for row in rows:
            values = [row.get(k, "") for k in keys]
            tree.insert("", tk.END, values=values)

    def _update_unassigned_tab_title(self) -> None:
        count = len(self.excel_unassigned_results)
        base_title = "Filas Excel no asignadas"
        title = f"{base_title} ({count})" if count != 0 else base_title
        self.tabs.tab(self.tab_unassigned_excel, text=title)

    def _on_main_selection_changed(self, _event) -> None:
        selected = self.tree_main.selection()
        if not selected:
            self.detail_var.set("Seleccione una fila en rojo para ver el detalle de diferencias (JP/JG).")
            return

        row = self.main_row_by_iid.get(selected[0])
        if not row:
            self.detail_var.set("Seleccione una fila en rojo para ver el detalle de diferencias (JP/JG).")
            return

        details = row.get("mismatch_details", [])
        if not details:
            self.detail_var.set("La fila seleccionada no tiene discordancias detalladas.")
            return

        lines: list[str] = []
        for mm in details[:2]:
            campo = str(mm.get("campo", ""))
            jp = str(mm.get("pdf", "") or "")
            jg = str(mm.get("excel", "") or "")
            lines.append(f"{campo}: [JP] {jp}")
            lines.append(f"{campo}: [JG] {jg}")

        self.detail_var.set("\n".join(lines))

    def _on_main_double_click(self, event) -> None:
        row_id = self.tree_main.identify_row(event.y)
        if not row_id:
            return

        row = self.main_row_by_iid.get(row_id)
        if not row:
            return

        source_name = str(row.get("source", "")).strip()
        page_number = int(row.get("page", 0))
        if not source_name or page_number <= 0:
            return

        row_key = self._build_row_key(source_name, page_number)
        current_manual = self.manual_jg_by_row_key.get(row_key, "")
        user_value = simpledialog.askstring(
            "Asignacion manual de JG",
            (
                f"Archivo {source_name}, pagina {page_number}.\n"
                "Indique el JG manual (ej: 12345-26).\n"
                "Puede dejarlo vacio para limpiar la asignacion."
            ),
            initialvalue=current_manual,
            parent=self,
        )

        if user_value is None:
            return

        normalized = self._normalize_manual_jg_input(user_value)
        if normalized:
            self.manual_jg_by_row_key[row_key] = normalized
        elif row_key in self.manual_jg_by_row_key:
            del self.manual_jg_by_row_key[row_key]

        self._apply_manual_jg_to_results()
        self._rebuild_excel_unassigned_results()
        self._refresh_tables()
        self._update_status_summary(prefix="Asignacion manual actualizada. ")

    def _on_main_click_action(self, event) -> None:
        row_id = self.tree_main.identify_row(event.y)
        if not row_id:
            return

        col_ref = self.tree_main.identify_column(event.x)
        if not col_ref.startswith("#"):
            return

        try:
            col_idx = int(col_ref[1:]) - 1
        except ValueError:
            return

        columns = list(self.tree_main["columns"])
        if col_idx < 0 or col_idx >= len(columns):
            return

        if columns[col_idx] != "jg_apply":
            return

        row = self.main_row_by_iid.get(row_id)
        if not row:
            return

        candidates_text = str(row.get("jg_candidates", "")).strip()
        if not candidates_text:
            return

        chosen = candidates_text.split("|")[0].strip()
        if not chosen:
            return

        source_name = str(row.get("source", "")).strip()
        page_number = int(row.get("page", 0))
        if not source_name or page_number <= 0:
            return

        row_key = self._build_row_key(source_name, page_number)
        self.manual_jg_by_row_key[row_key] = chosen
        self._apply_manual_jg_to_results()
        self._rebuild_excel_unassigned_results()
        self._refresh_tables()
        self._update_status_summary(prefix="JG candidato copiado a JG asignado. ")

    def _apply_manual_jg_to_results(self) -> None:
        manual_map = self.manual_jg_by_row_key

        for row in self.main_results:
            source_name = str(row.get("source", "")).strip()
            page_number = int(row.get("page", 0))
            row_key = self._build_row_key(source_name, page_number)
            if row_key in manual_map:
                row["jg"] = manual_map[row_key]
            else:
                row["jg"] = str(row.get("jg_display_base", "")).strip()

    def _rebuild_excel_unassigned_results(self) -> None:
        assigned_jg = {
            self._jg_number_only(self._get_effective_jg(row)).upper()
            for row in self.main_results
            if self._get_effective_jg(row)
        }

        unassigned: list[dict] = []
        for row in self.excel_data:
            row_jg = self._jg_number_only(self._format_jg(row.jg)).strip().upper()
            if row_jg in assigned_jg:
                continue

            importe_neto = self._format_decimal(row.importe_neto) if row.importe_neto is not None else ""
            unassigned.append(
                {
                    "jg": self._jg_number_only(self._format_jg(row.jg)),
                    "fecha": row.fecha,
                    "importe_neto": importe_neto,
                    "factura": row.factura,
                    "beneficiario": row.beneficiario,
                    "concepto": row.concepto,
                }
            )

        self.excel_unassigned_results = unassigned

    def _count_found(self) -> int:
        return sum(1 for row in self.main_results if self._get_effective_jg(row))

    def _count_status(self, status: str) -> int:
        return sum(1 for row in self.main_results if row.get("match_status") == status)

    def _count_empty_assigned_jg(self) -> int:
        return sum(1 for row in self.main_results if not str(row.get("jg", "")).strip())

    def _update_status_summary(self, prefix: str = "") -> None:
        self.status_var.set(
            (
                f"{prefix}"
                f"Encontrados: {self._count_found()} | "
                f"JG asignado vacio: {self._count_empty_assigned_jg()} | "
                f"No encontrados: {self._count_status('not_found')} | "
                f"Dudosos: {self._count_status('dubious')} | "
                f"Excel no asignadas: {len(self.excel_unassigned_results)}"
            )
        )

    def _get_duplicated_assigned_jg_set(self) -> set[str]:
        counts: dict[str, int] = {}
        for row in self.main_results:
            jg = str(row.get("jg", "")).strip().upper()
            if not jg:
                continue
            counts[jg] = counts.get(jg, 0) + 1
        return {jg for jg, count in counts.items() if count > 1}

    def _get_effective_jg(self, row: dict) -> str:
        source_name = str(row.get("source", "")).strip()
        page_number = int(row.get("page", 0))
        row_key = self._build_row_key(source_name, page_number)
        manual = str(self.manual_jg_by_row_key.get(row_key, "")).strip()
        if manual:
            return manual

        return str(row.get("jg_base", "")).strip()

    def reset_state(self) -> None:
        self.origin_pdf_paths = []
        self.jp_pdf_path = None
        self.excel_path = None

        self.jp_pdf_data = []
        self.excel_data = []
        self.main_results = []
        self.excel_unassigned_results = []
        self.manual_jg_by_row_key = {}

        self.lbl_drop_origin.configure(text="Arrastra aqui los PDF de DC, CJ, JG, etc.")
        self.lbl_drop_jp.configure(text="Arrastra aqui el PDF de Justif. Pago")
        self.lbl_drop_excel.configure(text="Arrastra aqui el Excel de Justif. Gastos")

        self.field_vars["anio"].set("")
        self.field_vars["dc"].set("")
        self.field_vars["cj"].set("")
        self.field_vars["jg_origen"].set("X")
        self.field_vars["rt"].set("—")
        self.field_vars["otros"].set("—")
        self.no_jp_mode_var.set(False)
        self.btn_load_jp.configure(state=tk.NORMAL)
        self.btn_load_excel.configure(state=tk.NORMAL)

        self.entry_map["dc"].configure(state="readonly")
        self.entry_map["cj"].configure(state="readonly")
        self.entry_map["rt"].configure(state="readonly")

        self._refresh_special_field_style("dc")
        self._refresh_special_field_style("cj")
        self._refresh_special_field_style("jg_origen")

        for item in self.tree_main.get_children():
            self.tree_main.delete(item)
        for item in self.tree_unassigned_excel.get_children():
            self.tree_unassigned_excel.delete(item)

        self.main_row_by_iid = {}
        self._update_unassigned_tab_title()
        self.detail_var.set("Seleccione una fila en rojo para ver el detalle de diferencias (JP/JG).")
        self.status_var.set("Estado reiniciado. Cargue PDF origen, JP y Excel para comenzar.")
        self._update_action_buttons()

    def confirm_reset_state(self) -> None:
        proceed = messagebox.askyesno(
            "Confirmar reinicio",
            "Se perderan los datos cargados y los resultados actuales. Desea continuar?",
        )
        if not proceed:
            return

        self.reset_state()

    def _read_pdf_data(self, pdf_path: Path) -> list[PdfPageData]:
        reader = PdfReader(str(pdf_path))
        parsed_rows: list[PdfPageData] = []

        for idx, page in enumerate(reader.pages):
            raw_text = page.extract_text() or ""
            parsed = self._parse_pdf_page(raw_text, idx + 1, pdf_path.name)
            parsed_rows.append(parsed)

        return parsed_rows

    def _parse_pdf_page(self, text: str, page_number: int, source_name: str) -> PdfPageData:
        text_single_line = self._sanitize_pdf_text(text)

        importe_match = re.search(r"Importe ordenado:\s*([0-9\.,]+)\s*EUR", text_single_line, flags=re.IGNORECASE)
        fecha_match = re.search(
            r"Fecha operaci[oó]n:\s*([0-9]{2}-[0-9]{2}-[0-9]{4})",
            text_single_line,
            flags=re.IGNORECASE,
        )

        beneficiario_match = re.search(
            r"EUR>>\s*(.*?)\s*POR CUENTA DE:",
            text_single_line,
            flags=re.IGNORECASE,
        )

        concepto_block_match = re.search(r"CONCEPTO:\s*(.*)$", text_single_line, flags=re.IGNORECASE)

        importe = self._parse_decimal(importe_match.group(1)) if importe_match else None
        fecha = fecha_match.group(1) if fecha_match else ""
        beneficiario = beneficiario_match.group(1).strip() if beneficiario_match else ""

        if concepto_block_match:
            concepto_block = concepto_block_match.group(1).strip()
            factura, concepto = self._split_factura_concepto(concepto_block)
        else:
            factura = ""
            concepto = ""

        return PdfPageData(
            source_name=source_name,
            page_number=page_number,
            fecha=fecha,
            importe=importe,
            factura=factura,
            beneficiario=beneficiario,
            concepto=concepto,
        )

    def _split_factura_concepto(self, concepto_block: str) -> tuple[str, str]:
        text = str(concepto_block or "").strip()
        factura_payload_match = re.match(r"Factura:\s*(.*)$", text, flags=re.IGNORECASE)
        if not factura_payload_match:
            return "", text

        payload = factura_payload_match.group(1).strip()

        for sep in re.finditer(r"\s-\s", payload):
            left = payload[: sep.start()].strip()
            right = payload[sep.end() :].strip()
            if not left or not right:
                continue
            if self._starts_with_four_letters(right):
                return left, right

        classic = re.match(r"^(.*?)\s-\s(.*)$", payload)
        if classic:
            return classic.group(1).strip(), classic.group(2).strip()

        return "", text

    def _starts_with_four_letters(self, text: str) -> bool:
        normalized = self._normalize_text(str(text or ""))
        head = normalized[:4]
        return len(head) == 4 and head.isalpha()

    @staticmethod
    def _sanitize_pdf_text(text: str) -> str:
        text = text.replace("\r", " ").replace("\n", " ")
        text = re.sub(r"\s+", " ", text)
        return text.strip()

    def _read_excel(self, excel_path: Path) -> list[ExcelRowData]:
        wb = load_workbook(filename=str(excel_path), data_only=True)
        ws = wb.active

        header_row_idx = self._locate_header_row(ws)
        col_index = self._map_columns(ws, header_row_idx)

        rows: list[ExcelRowData] = []
        for row_idx in range(header_row_idx + 1, ws.max_row + 1):
            jg_raw = ws.cell(row=row_idx, column=col_index["jg"]).value
            fecha_raw = ws.cell(row=row_idx, column=col_index["fecha"]).value
            total_raw = ws.cell(row=row_idx, column=col_index["importe_total"]).value
            ret_raw = ws.cell(row=row_idx, column=col_index["importe_retencion"]).value
            factura_raw = ws.cell(row=row_idx, column=col_index["factura"]).value
            beneficiario_raw = ws.cell(row=row_idx, column=col_index["beneficiario"]).value
            concepto_raw = ws.cell(row=row_idx, column=col_index["concepto"]).value

            fecha = self._normalize_excel_date(fecha_raw)
            total = self._parse_decimal(total_raw)
            ret = self._parse_decimal(ret_raw)
            factura = self._normalize_text(str(factura_raw or ""))
            beneficiario = self._normalize_text(str(beneficiario_raw or ""))
            concepto = self._normalize_text(str(concepto_raw or ""))

            if not self._is_valid_excel_row(
                jg_raw=jg_raw,
                fecha=fecha,
                importe_total=total,
                importe_retencion=ret,
                beneficiario=beneficiario,
                concepto=concepto,
            ):
                continue

            rows.append(
                ExcelRowData(
                    row_number=row_idx,
                    jg=str(jg_raw).strip(),
                    fecha=fecha,
                    importe_total=total,
                    importe_retencion=ret,
                    factura=factura,
                    beneficiario=beneficiario,
                    concepto=concepto,
                )
            )

        return rows

    @staticmethod
    def _locate_header_row(ws) -> int:
        for row_idx in range(1, ws.max_row + 1):
            row_values = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]
            normalized = [App._normalize_header_name(v) for v in row_values]
            if "n reg. justific." in normalized and "fecha pago" in normalized:
                return row_idx

        raise ValueError("No se encontro la fila de cabeceras en el Excel.")

    @staticmethod
    def _map_columns(ws, header_row_idx: int) -> dict[str, int]:
        headers = {}
        for col_idx in range(1, ws.max_column + 1):
            value = ws.cell(row=header_row_idx, column=col_idx).value
            name = App._normalize_header_name(value)
            headers[name] = col_idx

        required = {
            "jg": "n reg. justific.",
            "fecha": "fecha pago",
            "importe_total": "importe total",
            "importe_retencion": "importe retencion",
            "factura": "factura proveedor",
            "beneficiario": "nombre proveedor",
            "concepto": "descripcion",
        }

        mapped: dict[str, int] = {}
        for key, header in required.items():
            if header not in headers:
                raise ValueError(f"Falta la columna requerida en Excel: {header}")
            mapped[key] = headers[header]

        return mapped

    def _find_excel_matches(self, page: PdfPageData, excel_rows: list[ExcelRowData]) -> list[ExcelRowData]:
        page_fecha = self._normalize_text(page.fecha)
        page_factura_sig = self._factura_signature(page.factura)
        page_beneficiario_sig = self._beneficiary_signature(page.beneficiario)
        page_concepto_80 = self._concept_signature(page.concepto)

        matches: list[ExcelRowData] = []

        for row in excel_rows:
            if page_fecha != self._normalize_text(row.fecha):
                continue

            if page.importe is None or row.importe_neto is None:
                continue

            if page.importe.quantize(Decimal("0.01")) != row.importe_neto:
                continue

            if page_factura_sig != self._factura_signature(row.factura):
                continue

            if page_beneficiario_sig != self._beneficiary_signature(row.beneficiario):
                continue

            if page_concepto_80 != self._concept_signature(row.concepto):
                continue

            matches.append(row)

        return matches

    def _find_near_excel_candidates(self, page: PdfPageData, excel_rows: list[ExcelRowData], max_mismatches: int) -> list[dict]:
        evaluated: list[dict] = []
        page_fecha = self._normalize_text(page.fecha)
        page_factura_sig = self._factura_signature(page.factura)
        page_beneficiario_sig = self._beneficiary_signature(page.beneficiario)
        page_concept_sig = self._concept_signature(page.concepto)

        for row in excel_rows:
            mismatches: list[dict[str, str]] = []

            if page_fecha != self._normalize_text(row.fecha):
                mismatches.append({"campo": "Fecha", "pdf": page.fecha, "excel": row.fecha})

            page_importe = page.importe.quantize(Decimal("0.01")) if page.importe is not None else None
            row_importe = row.importe_neto
            if page_importe is None or row_importe is None or page_importe != row_importe:
                mismatches.append(
                    {
                        "campo": "Importe",
                        "pdf": self._format_decimal(page.importe),
                        "excel": self._format_decimal(row.importe_neto),
                    }
                )

            if page_factura_sig != self._factura_signature(row.factura):
                mismatches.append({"campo": "Factura", "pdf": page.factura, "excel": row.factura})

            row_benef_sig = self._beneficiary_signature(row.beneficiario)
            if page_beneficiario_sig != row_benef_sig:
                mismatches.append(
                    {
                        "campo": "Beneficiario",
                        "pdf": page.beneficiario,
                        "excel": row.beneficiario,
                    }
                )

            row_concept_sig = self._concept_signature(row.concepto)
            if page_concept_sig != row_concept_sig:
                mismatches.append({"campo": "Concepto", "pdf": page.concepto, "excel": row.concepto})

            if 0 < len(mismatches) <= max_mismatches:
                evaluated.append(
                    {
                        "row": row,
                        "mismatches": mismatches,
                        "mismatches_count": len(mismatches),
                    }
                )

        evaluated.sort(key=lambda x: (x["mismatches_count"], x["row"].row_number))
        return evaluated

    def _extract_probable_jg(self, page: PdfPageData) -> tuple[str, str]:
        year_match = re.match(r"^\d{2}-\d{2}-(\d{4})$", str(page.fecha or "").strip())
        if not year_match:
            return "", ""

        year2 = year_match.group(1)[-2:]

        factura_text = re.sub(r"\s+", "", str(page.factura or ""))
        factura_match = re.search(rf"(\d+)-{year2}\b", factura_text)
        if factura_match:
            number = factura_match.group(1).lstrip("0") or "0"
            return f"{number}-{year2}", "factura"

        concept_head = str(page.concepto or "")[:30]
        concept_head_no_spaces = re.sub(r"\s+", "", concept_head)
        concept_match = re.search(rf"\((\d+)-{year2}\)", concept_head_no_spaces)
        if concept_match:
            number = concept_match.group(1).lstrip("0") or "0"
            return f"{number}-{year2}", "concepto"

        return "", ""

    @staticmethod
    def _normalize_excel_date(value) -> str:
        if value is None:
            return ""

        if isinstance(value, datetime):
            return value.strftime("%d-%m-%Y")

        text = str(value).strip()
        if not text:
            return ""

        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
            try:
                dt = datetime.strptime(text, fmt)
                return dt.strftime("%d-%m-%Y")
            except ValueError:
                continue

        return text.replace("/", "-")

    @staticmethod
    def _parse_decimal(value) -> Decimal | None:
        if value is None:
            return None

        text = str(value).strip()
        if not text:
            return None

        text = text.replace("EUR", "").replace(" ", "")
        text = text.replace(".", "")
        text = text.replace(",", ".")

        try:
            return Decimal(text).quantize(Decimal("0.01"))
        except (InvalidOperation, ValueError):
            return None

    @staticmethod
    def _format_decimal(value: Decimal | None) -> str:
        if value is None:
            return ""
        return f"{value:.2f}".replace(".", ",")

    @staticmethod
    def _normalize_text(value: str) -> str:
        value = value.strip().upper()
        value = "".join(
            c for c in unicodedata.normalize("NFD", value) if unicodedata.category(c) != "Mn"
        )
        value = re.sub(r"\s+", " ", value)
        return value

    def _normalize_pdf_concept(self, value: str) -> str:
        text = self._normalize_text(value)
        previous = None
        while previous != text:
            previous = text
            text = re.sub(r"\b([A-Z])\s+([A-Z]{2,})\b", r"\1\2", text)
            text = re.sub(r"\b([A-Z]{2,})\s+([A-Z])\b", r"\1\2", text)
        return text

    def _concept_signature(self, value: str) -> str:
        text = self._normalize_pdf_concept(str(value or ""))
        text = re.sub(r"[^A-Z0-9]", "", text)
        return text[:75]

    def _beneficiary_signature(self, value: str) -> str:
        text = self._normalize_text(str(value or ""))
        text = text.replace("&", "")
        text = re.sub(r"\s+", "", text)
        return text[:25]

    def _factura_signature(self, value: str) -> str:
        text = self._normalize_text(str(value or ""))
        text = text.replace("_", "")
        text = re.sub(r"\s+", "", text)
        return text

    @staticmethod
    def _format_jg(raw_jg: str) -> str:
        text = str(raw_jg or "").strip()
        match = re.match(r"^(\d{4})/(\d+)$", text)
        if not match:
            return text

        year4 = match.group(1)
        numeric = match.group(2).lstrip("0") or "0"
        return f"{numeric}-{year4[-2:]}"

    def _normalize_manual_jg_input(self, value: str) -> str:
        text = str(value or "").strip().upper()
        if not text:
            return ""

        if re.match(r"^\d{4}/\d+$", text):
            return self._format_jg(text)

        direct = re.match(r"^(\d+)-(\d{2})$", text)
        if direct:
            number = direct.group(1).lstrip("0") or "0"
            return number

        return self._jg_number_only(text)

    @staticmethod
    def _jg_number_only(value: str) -> str:
        text = str(value or "").strip().upper()
        if not text:
            return ""

        match = re.match(r"^(\d+)\s*-\s*\d{2}$", text)
        if match:
            return match.group(1).lstrip("0") or "0"

        if text.isdigit():
            return text.lstrip("0") or "0"

        plain_match = re.match(r"^(\d+)", text)
        if plain_match:
            return plain_match.group(1).lstrip("0") or "0"

        return text

    @staticmethod
    def _normalize_header_name(value) -> str:
        text = str(value or "").strip().upper()
        text = "".join(
            c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn"
        )
        text = text.replace("N.", "N ")
        text = re.sub(r"N[^A-Z0-9 ]", "N ", text)
        text = re.sub(r"\s+", " ", text)
        return text.lower()

    @staticmethod
    def _is_valid_excel_row(
        jg_raw,
        fecha: str,
        importe_total: Decimal | None,
        importe_retencion: Decimal | None,
        beneficiario: str,
        concepto: str,
    ) -> bool:
        jg_text = str(jg_raw or "").strip()
        if not re.match(r"^20\d{2}/0{5}\d{5,}$", jg_text):
            return False

        if not re.match(r"^\d{2}-\d{2}-\d{4}$", fecha):
            return False

        if importe_total is None or importe_retencion is None:
            return False

        if not beneficiario:
            return False

        if not concepto:
            return False

        return True


if __name__ == "__main__":
    app = App()
    app.mainloop()
